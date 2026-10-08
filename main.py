"""
=========================================================
Local Claims Extractor (zip / folder / .txt input)
=========================================================
Reads every screenshot OCR .txt file from a batch ZIP (or a folder, or
single .txt files), cleans the OCR text, extracts the claim / ticket /
member fields with screen-aware regexes, and writes ONE Excel workbook.

Usage:
    python main.py "C:\\path\\to\\BATCH_xxx.zip"
    python main.py "C:\\path\\to\\BATCH_xxx.zip" -o output\\result.xlsx
    python main.py                      (prompts for the path)

Output (next to the input unless -o is given):
    <name>_extracted.xlsx   -- the workbook:
        Claim Summary         one row per claim, every field consolidated
        Screens               one row per screenshot, all fields + extras
        Dataset A - Claims / Dataset B - State Log / Scenario Summary /
        Consolidated / Raw Extracted / Data Dictionary
                              (the existing dashboard sheets)
    <name>_cleaned_text/    -- the cleaned OCR text of every file, for review
    knowledge_base.json     -- what the self-learning layer has learned so far
                               (agents, name variants, OCR/label fixes, new
                               screens/fields); reused and extended every run.

Works for any agent: the operator is detected from the batch itself (Welcome
banner, Assignee/Updated By, Rosco header), and new OCR patterns are learned
automatically -- see claims_extractor/pattern_learner.py.
"""

from __future__ import annotations

import argparse
import datetime as dt
import logging
import os
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from typing import Dict, List, Tuple

from claims_extractor import field_extractor
from claims_extractor.field_extractor import DocumentResult, build_dashboard_workbook
from claims_extractor.pattern_learner import KnowledgeBase, PatternLearner
from claims_extractor.screen_parser import extract_screen, learn_batch

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("claims_local")

# Batch file names carry the capture time in epoch milliseconds:
#   cid_<id>_cid_<id>_1784833939404_<uuid>.txt
_CAPTURE_MS_RE = re.compile(r"_(\d{13})_")


def _capture_time(name: str) -> str:
    m = _CAPTURE_MS_RE.search(name)
    if not m:
        return ""
    return dt.datetime.fromtimestamp(int(m.group(1)) / 1000).strftime("%Y-%m-%d %H:%M:%S")


def _sort_key(name: str):
    m = _CAPTURE_MS_RE.search(name)
    return (int(m.group(1)) if m else 0, name)


def read_inputs(path: str) -> List[Tuple[str, str]]:
    """Return [(file_name, raw_text)] in capture order."""
    p = Path(path)
    items: List[Tuple[str, str]] = []
    if p.is_file() and p.suffix.lower() == ".zip":
        with zipfile.ZipFile(p) as z:
            for n in z.namelist():
                if n.lower().endswith(".txt"):
                    items.append((n.split("/")[-1], z.read(n).decode("utf-8", errors="replace")))
    elif p.is_dir():
        for f in p.rglob("*.txt"):
            items.append((f.name, f.read_text(encoding="utf-8", errors="replace")))
    elif p.is_file() and p.suffix.lower() == ".txt":
        items.append((p.name, p.read_text(encoding="utf-8", errors="replace")))
    else:
        raise SystemExit(f"Not a .zip, folder or .txt file: {path}")
    return sorted(items, key=lambda it: _sort_key(it[0]))


# ---------------------------------------------------------------------------
# Extra sheets
# ---------------------------------------------------------------------------

def _all(r: DocumentResult) -> Dict[str, str]:
    d = {k: v for k, v in r.fields.items()}
    d.update(r.extra_fields)
    return {k: "; ".join(map(str, v)) if isinstance(v, list) else str(v) for k, v in d.items()}


def build_screens_rows(results: List[Tuple[str, DocumentResult]], times: Dict[str, str]) -> List[Dict]:
    preferred = [
        "Agent", "Claim ID", "Ticket ID", "OSC ID", "Member ID", "Member Name", "Claim State",
        "Blocker", "Ticket State",
    ]
    rows = []
    for seq, (name, r) in enumerate(results, start=1):
        vals = _all(r)
        row = {"Seq": seq, "Capture Time": times.get(name, ""), "Screen Type": r.detected_application}
        for k in preferred:
            row[k] = vals.pop(k, "")
        row.update(dict(sorted(vals.items())))
        row["Flags"] = "; ".join(f"{k}: {v}" for k, v in r.flags.items())
        row["Source File"] = name
        rows.append(row)
    return rows


def _mode(values: List[str]) -> str:
    values = [v for v in values if v]
    return Counter(values).most_common(1)[0][0] if values else ""


def _history(values: List[str]) -> str:
    return " -> ".join(dict.fromkeys(v for v in values if v))


def build_claim_summary(results: List[Tuple[str, DocumentResult]], times: Dict[str, str]) -> List[Dict]:
    """One row per Claim ID. Stable attributes take the most common value
    across that claim's screenshots (outvotes a one-off OCR misread);
    states take the LATEST value, with the full history alongside."""
    by_claim: Dict[str, List[Tuple[str, Dict[str, str]]]] = {}
    for name, r in results:
        cid = r.fields.get("Claim ID")
        if isinstance(cid, str) and cid:
            by_claim.setdefault(cid, []).append((name, _all(r)))

    rows = []
    for cid, shots in by_claim.items():
        col = lambda k: [v.get(k, "") for _n, v in shots]  # noqa: E731
        last = lambda k: next((x for x in reversed(col(k)) if x), "")  # noqa: E731
        names = [n for n in col("Member Name") if n]
        rows.append({
            "Agent": _mode(col("Agent")),
            "Claim ID": cid,
            "Ticket ID": _mode(col("Ticket ID")),
            "Ticket Name": _mode(col("Ticket Name")),
            "Member Name": max(names, key=len) if names else "",
            "OSC ID": _mode(col("OSC ID")),
            "Member ID": _mode(col("Member ID")),
            "Member State": _mode(col("Member State")),
            "Claim Type": _mode(col("Claim Type")),
            "Claim State (latest)": last("Claim State"),
            "Claim State History": _history(col("Claim State")),
            "Blocker (latest)": last("Blocker"),
            "Blocker History": _history(col("Blocker")),
            "Ticket State (latest)": last("Ticket State"),
            "Ticket State History": _history(col("Ticket State")),
            "Skill/Code": _mode(col("Code")) or _mode(col("Skill")),
            "Billed": _mode(col("Billed")),
            "Allowed": _mode(col("Allowed")),
            "Allowed %": _mode(col("Allowed %")),
            "Pay provider": _mode(col("Pay provider")),
            "Pay member": _mode(col("Pay member")),
            "Member responsibility": _mode(col("Member responsibilty")),
            "Interest Accrued": _mode(col("Interest Accrued")),
            "Service Start Date": _mode(col("Service Start Date")),
            "Service End Date": _mode(col("Service End Date")),
            "Received date": _mode(col("Received date")),
            "Prompt Process Due Date": _mode(col("Prompt Process Due Date")),
            "Provider Claim#": _mode(col("Provider Claim#")),
            "Transmission ID": _mode(col("Transmission ID")),
            "Billing Type": _mode(col("Billing Type")),
            "Stay type": _mode(col("Stay type")),
            "Place of Service": _mode(col("Place of Service")),
            "DRG": _mode(col("DRG")),
            "Goal": _mode(col("Goal")),
            "Authorization Required": _mode(col("Authorization Required")),
            "Medicaid Reclamation": _mode(col("Medicaid Reclamation")),
            "Submission Type": _mode(col("Submission Type")),
            "Vendor Source": _mode(col("Vendor Source")),
            "Payment Authority": _mode(col("Payment Authority")),
            "Tenant": _mode(col("Tenant")),
            "Assignee": _mode(col("Assignee")),
            "Ticket Due": _mode(col("Due")),
            "SLA": _mode(col("SLA")),
            "Ticket Created On": _mode(col("Created On")),
            "Ticket Updated On (latest)": last("Updated On"),
            "Ticket Resolved On": last("Resolved On"),
            "Resolution": last("Resolution"),
            "Screenshots": len(shots),
            "First Capture": times.get(shots[0][0], ""),
            "Last Capture": times.get(shots[-1][0], ""),
        })
    return rows


def write_extra_sheets(xlsx: str, summary: List[Dict], screens: List[Dict],
                       learned: Dict[str, List[Dict]] = None) -> None:
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(xlsx, engine="openpyxl", mode="a") as w:
        pd.DataFrame(summary).to_excel(w, index=False, sheet_name="Claim Summary")
        pd.DataFrame(screens).to_excel(w, index=False, sheet_name="Screens")
        if learned is not None:
            cols = ["Type", "Pattern", "Learned", "Seen", "Status", "Detail"]
            pd.DataFrame(learned["new"], columns=["Type", "Pattern", "Learned"]).to_excel(
                w, index=False, sheet_name="New This Batch")
            pd.DataFrame(learned["all"], columns=cols).to_excel(w, index=False, sheet_name="Learned Patterns")

    wb = load_workbook(xlsx)
    # Put the two new sheets first -- they are the ones to read.
    for title in ("Screens", "Claim Summary"):
        ws = wb[title]
        wb.move_sheet(ws, offset=-wb.index(ws))
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="FF1F4E78", end_color="FF1F4E78", fill_type="solid")
    flag_fill = PatternFill(start_color="FFFFF2CC", end_color="FFFFF2CC", fill_type="solid")
    for title in ("Claim Summary", "Screens", "New This Batch", "Learned Patterns"):
        if title not in wb.sheetnames:
            continue
        ws = wb[title]
        for c in ws[1]:
            c.font, c.fill = header_font, header_fill
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        flag_col = next((c.column for c in ws[1] if c.value == "Flags"), None)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial")
                if c.column == flag_col and c.value:
                    c.fill = flag_fill
        for col in ws.columns:
            longest = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(longest + 2, 10), 45)
    wb.active = 0
    wb.save(xlsx)


# ---------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description="Extract claim fields from a batch of OCR .txt files.")
    ap.add_argument("input", nargs="?", help="Batch .zip, a folder of .txt files, or one .txt file")
    ap.add_argument("-o", "--output", help="Output .xlsx path (default: next to the input)")
    ap.add_argument("--kb", help="Knowledge base JSON (default: knowledge_base.json next to main.py)")
    ap.add_argument("--no-learn", action="store_true",
                    help="Use the fixed rules only -- no learning, knowledge base not read or written")
    ap.add_argument("--push-sql", action="store_true",
                    help="Also upsert the results into Azure SQL (see push_to_sql.py / .env)")
    args = ap.parse_args()

    src = args.input or input("Path to the batch .zip (or folder): ").strip().strip('"')
    src_path = Path(src)
    stem = src_path.stem if src_path.suffix else src_path.name
    out = Path(args.output) if args.output else src_path.parent / f"{stem}_extracted.xlsx"
    clean_dir = out.parent / f"{stem}_cleaned_text"
    clean_dir.mkdir(parents=True, exist_ok=True)

    items = read_inputs(src)
    log.info("Found %d text file(s) in %s", len(items), src)

    learner = None
    if not args.no_learn:
        kb = KnowledgeBase.load(args.kb)
        learner = PatternLearner(kb)
        # Pass 1: learn this batch's agent, name variants and OCR misreads.
        learn_batch(learner, [raw for _n, raw in items], src_path.name)
        log.info("Agent detected: %s %s", learner.agent or "(none)", learner.agent_evidence or "")
        # Dataset A/B "User"/"Name" columns snap to the learned spellings.
        field_extractor._KNOWN_USER_NAMES = list(dict.fromkeys(
            ([learner.agent] if learner.agent else []) + list(kb["people"])))

    results: List[Tuple[str, DocumentResult]] = []
    times: Dict[str, str] = {}
    for name, raw in items:
        try:
            r, cleaned = extract_screen(raw, learner=learner, source_name=name)
        except Exception:
            log.exception("Failed on %s", name)
            continue
        (clean_dir / name).write_text(cleaned, encoding="utf-8")
        times[name] = _capture_time(name)
        results.append((name, r))

    if not results:
        log.warning("No documents extracted -- nothing to write.")
        return

    try:
        build_dashboard_workbook(results, str(out))
    except PermissionError:
        raise SystemExit(f"Cannot write {out} -- close it in Excel and run again.")
    summary = build_claim_summary(results, times)
    learned = None
    if learner:
        learner.finish_batch(len(results))
        learned = learner.report_rows()
        learner.kb.save()
    write_extra_sheets(str(out), summary, build_screens_rows(results, times), learned)

    apps = Counter(r.detected_application for _n, r in results)
    log.info("Screens by type: %s", dict(apps))
    log.info("Claims found: %s", ", ".join(row["Claim ID"] for row in summary))
    log.info("Workbook: %s", out.resolve())
    log.info("Cleaned text: %s", clean_dir.resolve())
    if learner:
        log.info("New patterns learned this batch: %d (see 'New This Batch' sheet)", len(learned["new"]))
        for row in learned["new"][:25]:
            log.info("  %-22s %s -> %s", row["Type"], row["Pattern"], row["Learned"])
        log.info("Knowledge base: %s", learner.kb.path.resolve())

    if args.push_sql:
        from claims_extractor.sql_client import push_results
        from push_to_sql import log_totals
        log_totals(push_results(results))


if __name__ == "__main__":
    main()
